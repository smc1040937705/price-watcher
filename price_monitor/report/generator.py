import os
from datetime import datetime
from typing import List, Dict, Any, Optional
import json

from .analyzer import ProductAnalysis, PriceAnalyzer, DataSourceInfo
from .visualizer import PriceVisualizer
from ..logger import setup_logger

class ReportGenerator:
    def __init__(self, output_dir: str = 'reports', db_path: str = None, json_path: str = None):
        self.logger = setup_logger(self.__class__.__name__)
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
        
        self.analyzer = PriceAnalyzer(db_path=db_path, json_path=json_path)
        self.visualizer = PriceVisualizer(output_dir=os.path.join(output_dir, 'charts'))
        
        self.analyses: List[ProductAnalysis] = []
        self.summary: Dict[str, Any] = {}
        self.charts: Dict[str, str] = {}
        self._base_filename: str = ""
    
    def prepare_data(self, use_json: bool = False, prefer_source: str = "auto"):
        self.analyzer.validate_data_sources()
        self.analyses = self.analyzer.analyze_all(use_json=use_json, prefer_source=prefer_source)
        self.summary = self.analyzer.get_summary(self.analyses)
        self._base_filename = f'price_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}'
        self.logger.info(f"已分析 {len(self.analyses)} 个商品")
        
        source_info = self.summary.get('data_source_info')
        if source_info and source_info.has_conflict:
            self.logger.warning(
                f"数据源存在冲突: {len(source_info.conflict_details)} 个问题，"
                f"报告将标注数据来源"
            )
    
    def generate_charts(self, top_n: int = 5):
        if not self.analyses:
            self.logger.warning("没有分析数据，无法生成图表")
            return
        
        top_products = sorted(self.analyses, 
                             key=lambda x: abs(x.change_percent_recent), 
                             reverse=True)[:top_n]
        self.charts = self.visualizer.generate_all_charts(self.analyses, top_products)
    
    def _get_base_filename(self) -> str:
        if not self._base_filename:
            self._base_filename = f'price_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}'
        return self._base_filename
    
    def generate_markdown_report(self, filename: str = None) -> str:
        if not self.analyses:
            self.prepare_data()
        
        if filename is None:
            filename = f'{self._get_base_filename()}.md'
        
        filepath = os.path.join(self.output_dir, filename)
        
        md_content = self._build_markdown_content()
        
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(md_content)
        
        self.logger.info(f"Markdown报告已生成: {filepath}")
        return filepath
    
    def _format_trend_table(self, trend, period_name: str) -> List[str]:
        if not trend:
            return [f"| {period_name} | - | - | - | - | - | - | - |"]
        
        return [f"| {period_name} | ¥{trend.min_price:.2f} | ¥{trend.max_price:.2f} | "
                f"¥{trend.avg_price:.2f} | ¥{trend.price_range:.2f} | "
                f"{trend.volatility:.2f}% | {trend.trend_direction} | "
                f"{trend.change_percent:+.2f}% |"]
    
    def _build_markdown_content(self) -> str:
        lines = []
        
        lines.append("# 📊 电商价格监控分析报告")
        lines.append("")
        lines.append(f"**报告生成时间**: {self.summary.get('generated_at', datetime.now().strftime('%Y-%m-%d %H:%M:%S'))}")
        
        source_info = self.summary.get('data_source_info')
        if source_info:
            lines.append(f"**数据源**: {self._get_source_description(source_info)}")
            if source_info.last_updated:
                lines.append(f"**最后更新**: {source_info.last_updated}")
        
        lines.append("")
        
        if source_info and source_info.has_conflict:
            lines.append("## ⚠️ 数据源冲突警告")
            lines.append("")
            lines.append("检测到以下数据不一致问题：")
            lines.append("")
            for conflict in source_info.conflict_details[:5]:
                lines.append(f"- {conflict}")
            if len(source_info.conflict_details) > 5:
                lines.append(f"- ... 还有 {len(source_info.conflict_details) - 5} 个问题")
            lines.append("")
            lines.append("> 建议：运行 `python cli.py export_data` 重新导出JSON数据以保持一致")
            lines.append("")
        
        lines.append("## 📈 监控概览")
        lines.append("")
        lines.append(f"| 指标 | 数值 |")
        lines.append("|:---|:---|")
        lines.append(f"| 监控商品总数 | {self.summary['total_products']} |")
        lines.append(f"| 降价商品数 | {self.summary['total_price_down']} |")
        lines.append(f"| 涨价商品数 | {self.summary['total_price_up']} |")
        lines.append(f"| 达到目标价商品 | {self.summary['total_target_reached']} |")
        lines.append(f"| 平均价格变动 | {self.summary['avg_change_percent']:.2f}% |")
        lines.append("")
        
        bargain_products = self.analyzer.get_bargain_products(self.analyses)
        if bargain_products:
            lines.append("## 📉 降价商品清单")
            lines.append("")
            lines.append("| 商品名称 | 平台 | 当前价格 | 降价幅度 | 目标价 | 数据源 |")
            lines.append("|:---|:---|:---:|:---:|:---:|:---:|")
            for a in sorted(bargain_products, key=lambda x: x.change_percent_recent):
                target_str = f"¥{a.target_price:.2f}" if a.target_price else "-"
                lines.append(f"| {a.product_name[:40]} | {a.platform} | ¥{a.current_price:.2f} | {a.change_percent_recent:.1f}% | {target_str} | {a.data_source} |")
            lines.append("")
        
        price_up_products = self.analyzer.get_price_up_products(self.analyses)
        if price_up_products:
            lines.append("## 📈 涨价商品清单")
            lines.append("")
            lines.append("| 商品名称 | 平台 | 当前价格 | 涨价幅度 | 目标价 | 数据源 |")
            lines.append("|:---|:---|:---:|:---:|:---:|:---:|")
            for a in sorted(price_up_products, key=lambda x: x.change_percent_recent, reverse=True):
                target_str = f"¥{a.target_price:.2f}" if a.target_price else "-"
                lines.append(f"| {a.product_name[:40]} | {a.platform} | ¥{a.current_price:.2f} | +{a.change_percent_recent:.1f}% | {target_str} | {a.data_source} |")
            lines.append("")
        
        target_reached = self.analyzer.get_target_reached_products(self.analyses)
        if target_reached:
            lines.append("## 🎯 达到目标价商品")
            lines.append("")
            lines.append("| 商品名称 | 当前价格 | 目标价 | 差价 | 数据源 |")
            lines.append("|:---|:---:|:---:|:---:|:---:|")
            for a in target_reached:
                diff = a.target_price - a.current_price
                lines.append(f"| {a.product_name[:40]} | ¥{a.current_price:.2f} | ¥{a.target_price:.2f} | ¥{diff:.2f} | {a.data_source} |")
            lines.append("")
        
        lines.append("## 📊 重点商品趋势分析")
        lines.append("")
        
        top_volatile = self.analyzer.get_top_volatile_products(self.analyses, 5)
        if top_volatile:
            lines.append("### 价格波动最大的商品")
            lines.append("")
            for a in top_volatile:
                lines.append(f"#### {a.product_name}")
                lines.append("")
                lines.append(f"- **平台**: {a.platform}")
                lines.append(f"- **当前价格**: ¥{a.current_price:.2f}")
                if a.target_price:
                    lines.append(f"- **目标价格**: ¥{a.target_price:.2f}")
                lines.append(f"- **数据源**: {a.data_source}")
                lines.append("")
                
                lines.append("**价格趋势对比**:")
                lines.append("")
                lines.append("| 时间段 | 最低价 | 最高价 | 平均价 | 波动范围 | 波动率 | 趋势 | 变动幅度 |")
                lines.append("|:---|:---:|:---:|:---:|:---:|:---:|:---|:---:|")
                
                lines.extend(self._format_trend_table(a.trend_7d, "7天"))
                lines.extend(self._format_trend_table(a.trend_30d, "30天"))
                lines.extend(self._format_trend_table(a.trend_all, "全部"))
                lines.append("")
                
                if a.trend_all:
                    lines.append(f"- **最低价日期**: {a.trend_all.lowest_date or 'N/A'}")
                    lines.append(f"- **最高价日期**: {a.trend_all.highest_date or 'N/A'}")
                lines.append(f"- **商品链接**: [{a.url}]({a.url})")
                lines.append("")
        
        lines.append("## 📋 全部商品价格汇总")
        lines.append("")
        lines.append("| 商品名称 | 平台 | 当前价格 | 最低价 | 最高价 | 波动率 | 趋势 | 数据源 |")
        lines.append("|:---|:---|:---:|:---:|:---:|:---:|:---|:---:|")
        for a in self.analyses:
            trend_str = a.trend_all.trend_direction if a.trend_all else "N/A"
            change_str = f"({a.trend_all.change_percent:+.1f}%)" if a.trend_all else ""
            min_price = f"¥{a.trend_all.min_price:.2f}" if a.trend_all else "-"
            max_price = f"¥{a.trend_all.max_price:.2f}" if a.trend_all else "-"
            volatility = f"{a.trend_all.volatility:.1f}%" if a.trend_all else "-"
            lines.append(f"| {a.product_name[:35]} | {a.platform} | ¥{a.current_price:.2f} | {min_price} | {max_price} | {volatility} | {trend_str} {change_str} | {a.data_source} |")
        lines.append("")
        
        lines.append("## 📈 多时段趋势对比表")
        lines.append("")
        lines.append("| 商品名称 | 7天趋势 | 7天变动 | 30天趋势 | 30天变动 | 全部趋势 | 全部变动 |")
        lines.append("|:---|:---|:---:|:---|:---:|:---|:---:|")
        for a in self.analyses:
            t7 = a.trend_7d
            t30 = a.trend_30d
            ta = a.trend_all
            
            t7_str = t7.trend_direction if t7 else "无数据"
            t7_change = f"{t7.change_percent:+.1f}%" if t7 else "-"
            t30_str = t30.trend_direction if t30 else "无数据"
            t30_change = f"{t30.change_percent:+.1f}%" if t30 else "-"
            ta_str = ta.trend_direction if ta else "无数据"
            ta_change = f"{ta.change_percent:+.1f}%" if ta else "-"
            
            lines.append(f"| {a.product_name[:30]} | {t7_str} | {t7_change} | {t30_str} | {t30_change} | {ta_str} | {ta_change} |")
        lines.append("")
        
        lines.append("---")
        lines.append("")
        lines.append("*本报告由电商价格监控系统自动生成*")
        
        return '\n'.join(lines)
    
    def _get_source_description(self, source_info: DataSourceInfo) -> str:
        if source_info.source_type == "both":
            return "数据库 + JSON"
        elif source_info.source_type == "json":
            return "JSON文件"
        elif source_info.source_type == "database":
            return "SQLite数据库"
        else:
            return "无数据源"
    
    def generate_excel_report(self, filename: str = None) -> str:
        if not self.analyses:
            self.prepare_data()
        
        if filename is None:
            filename = f'{self._get_base_filename()}.xlsx'
        
        filepath = os.path.join(self.output_dir, filename)
        
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            self.logger.error("请安装 openpyxl: pip install openpyxl")
            return ""
        
        wb = openpyxl.Workbook()
        
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        warning_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        center_align = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        ws_summary = wb.active
        ws_summary.title = "监控概览"
        
        ws_summary['A1'] = "电商价格监控分析报告"
        ws_summary['A1'].font = Font(bold=True, size=16)
        ws_summary.merge_cells('A1:D1')
        
        ws_summary['A3'] = "报告生成时间"
        ws_summary['B3'] = self.summary.get('generated_at', datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
        
        source_info = self.summary.get('data_source_info')
        if source_info:
            ws_summary['A4'] = "数据源"
            ws_summary['B4'] = self._get_source_description(source_info)
            if source_info.last_updated:
                ws_summary['A5'] = "最后更新"
                ws_summary['B5'] = source_info.last_updated
        
        summary_data = [
            ['指标', '数值'],
            ['监控商品总数', self.summary['total_products']],
            ['降价商品数', self.summary['total_price_down']],
            ['涨价商品数', self.summary['total_price_up']],
            ['达到目标价商品', self.summary['total_target_reached']],
            ['平均价格变动', f"{self.summary['avg_change_percent']:.2f}%"]
        ]
        
        start_row = 7
        if source_info and source_info.has_conflict:
            ws_summary.cell(row=start_row, column=1, value="⚠️ 数据源冲突警告").fill = warning_fill
            start_row += 1
            for i, conflict in enumerate(source_info.conflict_details[:3]):
                ws_summary.cell(row=start_row + i, column=1, value=f"  • {conflict}")
            start_row += 4
        
        for row_idx, row_data in enumerate(summary_data, start=start_row):
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws_summary.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = center_align
                if row_idx == start_row:
                    cell.fill = header_fill
                    cell.font = header_font
        
        ws_summary.column_dimensions['A'].width = 20
        ws_summary.column_dimensions['B'].width = 15
        
        ws_all = wb.create_sheet("全部商品")
        
        headers = ['商品名称', '平台', '当前价格', '最低价', '最高价', '平均价', 
                   '波动范围', '波动率', '趋势', '变动幅度', '目标价', '是否达到目标', '数据源']
        
        for col_idx, header in enumerate(headers, start=1):
            cell = ws_all.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        
        for row_idx, a in enumerate(self.analyses, start=2):
            row_data = [
                a.product_name,
                a.platform,
                a.current_price,
                a.trend_all.min_price if a.trend_all else 0,
                a.trend_all.max_price if a.trend_all else 0,
                a.trend_all.avg_price if a.trend_all else 0,
                a.trend_all.price_range if a.trend_all else 0,
                f"{a.trend_all.volatility:.2f}%" if a.trend_all else "0%",
                a.trend_all.trend_direction if a.trend_all else "N/A",
                f"{a.trend_all.change_percent:+.2f}%" if a.trend_all else "0%",
                a.target_price if a.target_price else "-",
                "是" if a.reached_target else "否",
                a.data_source
            ]
            
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws_all.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = center_align
                
                if col_idx == 10:
                    if a.trend_all and a.trend_all.change_percent < 0:
                        cell.font = Font(color="00AA00")
                    elif a.trend_all and a.trend_all.change_percent > 0:
                        cell.font = Font(color="FF0000")
        
        for col_idx in range(1, len(headers) + 1):
            ws_all.column_dimensions[get_column_letter(col_idx)].width = 15
        ws_all.column_dimensions['A'].width = 40
        
        ws_trend = wb.create_sheet("多时段趋势对比")
        
        trend_headers = ['商品名称', '7天趋势', '7天变动', '7天记录数', 
                         '30天趋势', '30天变动', '30天记录数',
                         '全部趋势', '全部变动', '全部记录数', '数据源']
        
        for col_idx, header in enumerate(trend_headers, start=1):
            cell = ws_trend.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
        
        for row_idx, a in enumerate(self.analyses, start=2):
            t7 = a.trend_7d
            t30 = a.trend_30d
            ta = a.trend_all
            
            row_data = [
                a.product_name,
                t7.trend_direction if t7 else "无数据",
                f"{t7.change_percent:+.2f}%" if t7 else "-",
                t7.record_count if t7 else 0,
                t30.trend_direction if t30 else "无数据",
                f"{t30.change_percent:+.2f}%" if t30 else "-",
                t30.record_count if t30 else 0,
                ta.trend_direction if ta else "无数据",
                f"{ta.change_percent:+.2f}%" if ta else "-",
                ta.record_count if ta else 0,
                a.data_source
            ]
            
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws_trend.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border
                cell.alignment = center_align
        
        for col_idx in range(1, len(trend_headers) + 1):
            ws_trend.column_dimensions[get_column_letter(col_idx)].width = 12
        ws_trend.column_dimensions['A'].width = 40
        
        bargain_products = self.analyzer.get_bargain_products(self.analyses)
        if bargain_products:
            ws_bargain = wb.create_sheet("降价商品")
            
            headers_bargain = ['商品名称', '平台', '当前价格', '降价幅度', '目标价', '商品链接', '数据源']
            for col_idx, header in enumerate(headers_bargain, start=1):
                cell = ws_bargain.cell(row=1, column=col_idx, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = thin_border
            
            for row_idx, a in enumerate(sorted(bargain_products, key=lambda x: x.change_percent_recent), start=2):
                row_data = [
                    a.product_name,
                    a.platform,
                    a.current_price,
                    f"{a.change_percent_recent:.1f}%",
                    a.target_price if a.target_price else "-",
                    a.url,
                    a.data_source
                ]
                for col_idx, value in enumerate(row_data, start=1):
                    cell = ws_bargain.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = thin_border
                    cell.alignment = center_align
            
            for col_idx in range(1, len(headers_bargain) + 1):
                ws_bargain.column_dimensions[get_column_letter(col_idx)].width = 15
            ws_bargain.column_dimensions['A'].width = 40
            ws_bargain.column_dimensions['F'].width = 50
        
        price_up_products = self.analyzer.get_price_up_products(self.analyses)
        if price_up_products:
            ws_up = wb.create_sheet("涨价商品")
            
            headers_up = ['商品名称', '平台', '当前价格', '涨价幅度', '目标价', '商品链接', '数据源']
            for col_idx, header in enumerate(headers_up, start=1):
                cell = ws_up.cell(row=1, column=col_idx, value=header)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align
                cell.border = thin_border
            
            for row_idx, a in enumerate(sorted(price_up_products, key=lambda x: x.change_percent_recent, reverse=True), start=2):
                row_data = [
                    a.product_name,
                    a.platform,
                    a.current_price,
                    f"+{a.change_percent_recent:.1f}%",
                    a.target_price if a.target_price else "-",
                    a.url,
                    a.data_source
                ]
                for col_idx, value in enumerate(row_data, start=1):
                    cell = ws_up.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = thin_border
                    cell.alignment = center_align
            
            for col_idx in range(1, len(headers_up) + 1):
                ws_up.column_dimensions[get_column_letter(col_idx)].width = 15
            ws_up.column_dimensions['A'].width = 40
            ws_up.column_dimensions['F'].width = 50
        
        wb.save(filepath)
        self.logger.info(f"Excel报告已生成: {filepath}")
        return filepath
    
    def generate_pdf_report(self, filename: str = None) -> str:
        if not self.analyses:
            self.prepare_data()
        
        if filename is None:
            filename = f'{self._get_base_filename()}.pdf'
        
        filepath = os.path.join(self.output_dir, filename)
        
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
        except ImportError:
            self.logger.error("请安装 reportlab: pip install reportlab")
            return ""
        
        doc = SimpleDocTemplate(filepath, pagesize=A4, 
                                rightMargin=1*cm, leftMargin=1*cm,
                                topMargin=1*cm, bottomMargin=1*cm)
        
        styles = getSampleStyleSheet()
        
        try:
            font_paths = [
                'C:/Windows/Fonts/simhei.ttf',
                'C:/Windows/Fonts/msyh.ttc',
                '/System/Library/Fonts/PingFang.ttc',
                '/usr/share/fonts/truetype/wqy/wqy-microhei.ttc',
            ]
            font_registered = False
            for font_path in font_paths:
                if os.path.exists(font_path):
                    try:
                        pdfmetrics.registerFont(TTFont('ChineseFont', font_path))
                        font_registered = True
                        break
                    except:
                        continue
            
            if font_registered:
                title_style = ParagraphStyle(
                    'ChineseTitle',
                    parent=styles['Heading1'],
                    fontName='ChineseFont',
                    fontSize=18,
                    spaceAfter=20
                )
                heading_style = ParagraphStyle(
                    'ChineseHeading',
                    parent=styles['Heading2'],
                    fontName='ChineseFont',
                    fontSize=14,
                    spaceAfter=10
                )
                normal_style = ParagraphStyle(
                    'ChineseNormal',
                    parent=styles['Normal'],
                    fontName='ChineseFont',
                    fontSize=10
                )
                warning_style = ParagraphStyle(
                    'ChineseWarning',
                    parent=styles['Normal'],
                    fontName='ChineseFont',
                    fontSize=10,
                    textColor=colors.red
                )
            else:
                title_style = styles['Heading1']
                heading_style = styles['Heading2']
                normal_style = styles['Normal']
                warning_style = styles['Normal']
        except:
            title_style = styles['Heading1']
            heading_style = styles['Heading2']
            normal_style = styles['Normal']
            warning_style = styles['Normal']
        
        elements = []
        
        elements.append(Paragraph("电商价格监控分析报告", title_style))
        elements.append(Paragraph(f"报告生成时间: {self.summary.get('generated_at', datetime.now().strftime('%Y-%m-%d %H:%M:%S'))}", normal_style))
        
        source_info = self.summary.get('data_source_info')
        if source_info:
            elements.append(Paragraph(f"数据源: {self._get_source_description(source_info)}", normal_style))
        
        elements.append(Spacer(1, 20))
        
        if source_info and source_info.has_conflict:
            elements.append(Paragraph("数据源冲突警告", heading_style))
            elements.append(Paragraph("检测到数据不一致，建议重新导出JSON数据", warning_style))
            elements.append(Spacer(1, 10))
        
        elements.append(Paragraph("监控概览", heading_style))
        
        summary_data = [
            ['指标', '数值'],
            ['监控商品总数', str(self.summary['total_products'])],
            ['降价商品数', str(self.summary['total_price_down'])],
            ['涨价商品数', str(self.summary['total_price_up'])],
            ['达到目标价商品', str(self.summary['total_target_reached'])],
            ['平均价格变动', f"{self.summary['avg_change_percent']:.2f}%"]
        ]
        
        summary_table = Table(summary_data, colWidths=[4*cm, 3*cm])
        summary_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'ChineseFont' if 'ChineseFont' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ]))
        elements.append(summary_table)
        elements.append(Spacer(1, 20))
        
        elements.append(Paragraph("多时段趋势对比", heading_style))
        
        trend_header = ['商品名称', '7天趋势', '7天变动', '30天趋势', '30天变动', '全部趋势', '全部变动']
        trend_data = [trend_header]
        
        for a in self.analyses:
            t7 = a.trend_7d
            t30 = a.trend_30d
            ta = a.trend_all
            
            row = [
                a.product_name[:20],
                t7.trend_direction if t7 else "-",
                f"{t7.change_percent:+.1f}%" if t7 else "-",
                t30.trend_direction if t30 else "-",
                f"{t30.change_percent:+.1f}%" if t30 else "-",
                ta.trend_direction if ta else "-",
                f"{ta.change_percent:+.1f}%" if ta else "-"
            ]
            trend_data.append(row)
        
        trend_table = Table(trend_data, colWidths=[4*cm, 2*cm, 2*cm, 2*cm, 2*cm, 2*cm, 2*cm])
        trend_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'ChineseFont' if 'ChineseFont' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 7),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#D6DCE4')]),
        ]))
        elements.append(trend_table)
        elements.append(Spacer(1, 20))
        
        bargain_products = self.analyzer.get_bargain_products(self.analyses)
        if bargain_products:
            elements.append(Paragraph("降价商品清单", heading_style))
            
            bargain_data = [['商品名称', '平台', '当前价格', '降价幅度', '数据源']]
            for a in sorted(bargain_products, key=lambda x: x.change_percent_recent)[:20]:
                bargain_data.append([
                    a.product_name[:25],
                    a.platform,
                    f"¥{a.current_price:.2f}",
                    f"{a.change_percent_recent:.1f}%",
                    a.data_source
                ])
            
            bargain_table = Table(bargain_data, colWidths=[5*cm, 1.5*cm, 2.5*cm, 2*cm, 1.5*cm])
            bargain_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#70AD47')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'ChineseFont' if 'ChineseFont' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#E2EFDA')]),
            ]))
            elements.append(bargain_table)
            elements.append(Spacer(1, 20))
        
        price_up_products = self.analyzer.get_price_up_products(self.analyses)
        if price_up_products:
            elements.append(Paragraph("涨价商品清单", heading_style))
            
            up_data = [['商品名称', '平台', '当前价格', '涨价幅度', '数据源']]
            for a in sorted(price_up_products, key=lambda x: x.change_percent_recent, reverse=True)[:20]:
                up_data.append([
                    a.product_name[:25],
                    a.platform,
                    f"¥{a.current_price:.2f}",
                    f"+{a.change_percent_recent:.1f}%",
                    a.data_source
                ])
            
            up_table = Table(up_data, colWidths=[5*cm, 1.5*cm, 2.5*cm, 2*cm, 1.5*cm])
            up_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#C00000')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'ChineseFont' if 'ChineseFont' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8CBAD')]),
            ]))
            elements.append(up_table)
            elements.append(Spacer(1, 20))
        
        elements.append(Paragraph("全部商品价格汇总", heading_style))
        
        all_data = [['商品名称', '当前价格', '最低价', '最高价', '波动率', '趋势', '数据源']]
        for a in self.analyses:
            trend_str = a.trend_all.trend_direction if a.trend_all else "N/A"
            all_data.append([
                a.product_name[:20],
                f"¥{a.current_price:.2f}",
                f"¥{a.trend_all.min_price:.2f}" if a.trend_all else "-",
                f"¥{a.trend_all.max_price:.2f}" if a.trend_all else "-",
                f"{a.trend_all.volatility:.1f}%" if a.trend_all else "-",
                trend_str,
                a.data_source
            ])
        
        all_table = Table(all_data, colWidths=[4*cm, 2.5*cm, 2.5*cm, 2.5*cm, 2*cm, 1.5*cm, 1.5*cm])
        all_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'ChineseFont' if 'ChineseFont' in pdfmetrics.getRegisteredFontNames() else 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 7),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#D6DCE4')]),
        ]))
        elements.append(all_table)
        
        doc.build(elements)
        self.logger.info(f"PDF报告已生成: {filepath}")
        return filepath
    
    def generate_all_reports(self, use_json: bool = False, prefer_source: str = "auto") -> Dict[str, str]:
        self.prepare_data(use_json=use_json, prefer_source=prefer_source)
        
        self.generate_charts()
        
        reports = {}
        
        reports['markdown'] = self.generate_markdown_report()
        reports['excel'] = self.generate_excel_report()
        reports['pdf'] = self.generate_pdf_report()
        
        self.logger.info(f"所有报告已生成完成，共 {len(reports)} 个文件")
        return reports
